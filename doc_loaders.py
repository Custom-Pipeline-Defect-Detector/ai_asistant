"""
doc_loaders.py

Extract text from many document types in a robust, debuggable way.

Upgrades:
- Smart Auto-OCR:
  - PDF: try text extract first; OCR only if likely scanned (low text + image-heavy)
  - Images: OCR always (if enabled)
- Per-file diagnostics (method, extracted chars, warnings/errors)
- Offline-friendly (hardcoded Poppler + Tesseract paths)
"""
from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import subprocess
import tempfile
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Callable, Dict, Iterable, List, Optional, Tuple

# ============================================================
# 🔒 HARD-CODED PATHS (WINDOWS OFFLINE OCR)
# ============================================================

POPPLER_BIN = r"D:\offline_doc_assistant\poppler\Library\bin"
TESSERACT_EXE = r"C:\Program Files\Tesseract-OCR\tesseract.exe"

# ----------------------------
# Text cleanup helpers
# ----------------------------

_WS_RE = re.compile(r"[ \t]+\n")
_NL_RE = re.compile(r"\n{3,}")


def clean_text(s: str) -> str:
    s = (s or "").replace("\r\n", "\n").replace("\r", "\n")
    s = _WS_RE.sub("\n", s)
    s = _NL_RE.sub("\n\n", s)
    return s.strip()


def load_txt_like(path: Path) -> str:
    try:
        return clean_text(path.read_text(encoding="utf-8", errors="ignore"))
    except Exception:
        return clean_text(path.read_text(errors="ignore"))


# ----------------------------
# Diagnostics
# ----------------------------

@dataclass
class LoadReportEntry:
    source: str
    ext: str
    ok: bool
    method: str
    chars: int
    warning: str = ""
    error: str = ""

    def to_dict(self) -> Dict[str, object]:
        return asdict(self)


# ----------------------------
# Optional helpers
# ----------------------------

_TESS_LANG_CACHE: Optional[str] = None


def _normalize_ocr_lang(lang: str) -> str:
    t = (lang or "").strip().lower()
    if not t or t in {"auto", "zh", "chinese", "cn"}:
        return _auto_tesseract_lang()

    if t in {"zh-cn", "zh-hans", "hans", "chs"}:
        return "eng+chi_sim"
    if t in {"zh-tw", "zh-hk", "zh-hant", "hant", "cht"}:
        return "eng+chi_tra"

    return (lang or "eng").strip()


def _auto_tesseract_lang() -> str:
    global _TESS_LANG_CACHE
    if _TESS_LANG_CACHE:
        return _TESS_LANG_CACHE

    best = "eng+chi_sim"
    try:
        import pytesseract  # type: ignore
        pytesseract.pytesseract.tesseract_cmd = TESSERACT_EXE

        langs = set()
        try:
            for x in pytesseract.get_languages(config=""):
                langs.add(str(x).strip())
        except Exception:
            langs = set()

        if "eng" not in langs:
            langs.add("eng")

        has_sim = "chi_sim" in langs
        has_tra = "chi_tra" in langs

        if has_sim and has_tra:
            best = "eng+chi_sim+chi_tra"
        elif has_sim:
            best = "eng+chi_sim"
        elif has_tra:
            best = "eng+chi_tra"
        else:
            best = "eng"

    except Exception:
        best = "eng"

    _TESS_LANG_CACHE = best
    return best


def _default_soffice() -> Optional[str]:
    env = os.environ.get("SOFFICE_PATH", "").strip().strip('"')
    if env:
        p = Path(env)
        if p.exists() and p.is_file():
            return str(p)

    for c in ("soffice", "soffice.exe"):
        p = shutil.which(c)
        if p:
            return p

    if os.name == "nt":
        candidates = [
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        ]
        for c in candidates:
            if Path(c).exists():
                return c

    return None


# ----------------------------
# Core loaders (direct)
# ----------------------------

def _load_pdf_text(path: Path) -> str:
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    parts: List[str] = []
    for page in reader.pages:
        try:
            t = page.extract_text() or ""
        except Exception:
            t = ""
        if t.strip():
            parts.append(t)
    return clean_text("\n".join(parts))


def _pdf_looks_scanned(
    path: Path,
    extracted_text: str,
    *,
    min_chars: int = 60,
    max_pages_check: int = 5,
    image_page_ratio_threshold: float = 0.6,
) -> bool:
    """
    Heuristic:
    - If extracted text is already decent -> not scanned
    - Else inspect first N pages for image XObjects; if most pages have images -> likely scanned
    """
    t = (extracted_text or "").strip()
    if len(t) >= min_chars:
        return False

    try:
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        n_pages = min(len(reader.pages), max_pages_check)
        if n_pages <= 0:
            return True

        img_pages = 0
        for i in range(n_pages):
            page = reader.pages[i]
            try:
                res = page.get("/Resources")
                if not res:
                    continue
                xobj = res.get("/XObject")
                if not xobj:
                    continue
                # XObject is a dict-like of objects; count /Subtype /Image
                has_image = False
                for _, obj in xobj.items():
                    try:
                        o = obj.get_object()
                        if o.get("/Subtype") == "/Image":
                            has_image = True
                            break
                    except Exception:
                        continue
                if has_image:
                    img_pages += 1
            except Exception:
                continue

        ratio = img_pages / max(1, n_pages)
        return ratio >= image_page_ratio_threshold
    except Exception:
        # If we can't inspect, fall back to low-text heuristic:
        return len(t) < min_chars


def _load_docx(path: Path) -> str:
    from docx import Document
    doc = Document(str(path))
    parts: List[str] = [p.text for p in doc.paragraphs if p.text.strip()]
    return clean_text("\n".join(parts))


def _load_html(path: Path) -> str:
    raw = load_txt_like(path)
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(raw, "html.parser")
    for tag in soup(["script", "style", "noscript"]):
        tag.extract()
    return clean_text(soup.get_text(separator="\n"))


def _load_csv(path: Path, delimiter: str = ",") -> str:
    import csv
    rows: List[str] = []
    with path.open("r", encoding="utf-8", errors="ignore", newline="") as f:
        reader = csv.reader(f, delimiter=delimiter)
        for row in reader:
            rows.append("\t".join([str(c).strip() for c in row]))
    return clean_text("\n".join(rows))


def _load_xlsx(path: Path) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    parts: List[str] = []
    for ws in wb.worksheets:
        parts.append(f"=== Sheet: {ws.title} ===")
        for row in ws.iter_rows(values_only=True):
            vals = [str(v).strip() for v in row if v is not None]
            if vals:
                parts.append("\t".join(vals))
    return clean_text("\n".join(parts))


def _load_pptx(path: Path) -> str:
    from pptx import Presentation
    prs = Presentation(str(path))
    parts: List[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        parts.append(f"=== Slide {i} ===")
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                parts.append(shape.text.strip())
    return clean_text("\n".join(parts))


def _load_rtf(path: Path) -> str:
    raw = load_txt_like(path)
    raw = re.sub(r"{\\[^}]+}", " ", raw)
    raw = re.sub(r"\\[a-zA-Z]+\d* ?", " ", raw)
    raw = raw.replace("{", " ").replace("}", " ")
    return clean_text(raw)


def _load_json(path: Path) -> str:
    raw = load_txt_like(path)
    try:
        obj = json.loads(raw)
        return clean_text(json.dumps(obj, ensure_ascii=False, indent=2))
    except Exception:
        return raw


# ----------------------------
# OCR (optional)
# ----------------------------

def _ocr_image_file(path: Path, lang: str = "auto") -> str:
    from PIL import Image  # type: ignore
    import pytesseract  # type: ignore

    pytesseract.pytesseract.tesseract_cmd = TESSERACT_EXE
    lang = _normalize_ocr_lang(lang)

    img = Image.open(str(path))
    return clean_text(pytesseract.image_to_string(img, lang=lang))


def _ocr_pdf(
    path: Path,
    lang: str = "auto",
    max_pages: Optional[int] = None,
    poppler_path: Optional[str] = None,
) -> str:
    from pdf2image import convert_from_path  # type: ignore
    import pytesseract  # type: ignore

    pytesseract.pytesseract.tesseract_cmd = TESSERACT_EXE
    if poppler_path is None:
        poppler_path = POPPLER_BIN

    lang = _normalize_ocr_lang(lang)

    kwargs = dict(
        dpi=250,
        first_page=1,
        poppler_path=poppler_path,
    )
    if max_pages is not None:
        kwargs["last_page"] = max_pages

    images = convert_from_path(str(path), **kwargs)

    texts = []
    for img in images:
        txt = pytesseract.image_to_string(img, lang=lang)
        if txt.strip():
            texts.append(txt)

    return clean_text("\n\n".join(texts))


# ----------------------------
# LibreOffice conversion (optional)
# ----------------------------

def _convert_with_libreoffice(
    soffice_path: str,
    input_path: Path,
    outdir: Path,
    target: str,
    timeout_s: int = 90,
) -> Path:
    outdir.mkdir(parents=True, exist_ok=True)
    cmd = [
        soffice_path,
        "--headless",
        "--nologo",
        "--nofirststartwizard",
        "--convert-to",
        target,
        "--outdir",
        str(outdir),
        str(input_path),
    ]
    proc = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout_s)
    if proc.returncode != 0:
        raise RuntimeError(proc.stderr.strip() or proc.stdout.strip())

    stem = input_path.stem
    expected = outdir / f"{stem}.txt"
    return expected


# ----------------------------
# Loader registry
# ----------------------------

LoaderFn = Callable[[Path], str]


def _text_like_exts() -> set:
    return {
        ".txt", ".md", ".mdx", ".log", ".ini", ".cfg", ".conf", ".env", ".xml",
        ".yml", ".yaml", ".tsv", ".tex",
        ".py", ".js", ".ts", ".tsx", ".jsx",
        ".java", ".kt", ".c", ".h", ".cpp", ".hpp", ".cs",
        ".go", ".rs", ".rb", ".php", ".sql", ".sh", ".bat", ".ps1",
        ".toml", ".gitignore",
        ".json", ".jsonl", ".ndjson",
        ".html", ".htm",
    }


DIRECT_LOADERS: Dict[str, LoaderFn] = {
    ".pdf": _load_pdf_text,
    ".docx": _load_docx,
    ".html": _load_html,
    ".htm": _load_html,
    ".csv": lambda p: _load_csv(p, delimiter=","),
    ".tsv": lambda p: _load_csv(p, delimiter="\t"),
    ".json": _load_json,
    ".rtf": _load_rtf,
    ".xlsx": _load_xlsx,
    ".xlsm": _load_xlsx,
    ".pptx": _load_pptx,
}

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff"}
CONVERTABLE_OFFICE_EXTS = {".doc", ".odt", ".xls", ".ods", ".ppt", ".odp"}


# ----------------------------
# Main entry
# ----------------------------

def load_documents(
    paths: Iterable[Path],
    *,
    enable_ocr: bool = True,
    ocr_lang: str = "auto",
    enable_libreoffice: bool = True,
    soffice_path: Optional[str] = None,
    # PDF auto-OCR heuristics
    scanned_pdf_min_chars: int = 60,
    pdf_ocr_max_pages: Optional[int] = None,
    poppler_path: Optional[str] = None,
) -> Tuple[List[Dict[str, str]], List[LoadReportEntry]]:

    docs: List[Dict[str, str]] = []
    report: List[LoadReportEntry] = []

    soffice = soffice_path or _default_soffice()
    ocr_lang = _normalize_ocr_lang(ocr_lang)

    for p in paths:
        p = Path(p)
        ext = p.suffix.lower()

        text = ""
        method = "none"
        warning = ""
        error = ""
        ok = False

        try:
            if ext in _text_like_exts():
                text = load_txt_like(p)
                method = "text"
                ok = bool(text.strip())

            elif ext in DIRECT_LOADERS:
                method = f"direct:{ext}"
                text = DIRECT_LOADERS[ext](p)
                text = clean_text(text)
                ok = bool(text.strip())

                # Smart PDF OCR (AUTO)
                if ext == ".pdf" and enable_ocr:
                    if _pdf_looks_scanned(p, text, min_chars=scanned_pdf_min_chars):
                        method = "ocr:pdf(auto)"
                        text = _ocr_pdf(
                            p,
                            lang=ocr_lang,
                            max_pages=pdf_ocr_max_pages,
                            poppler_path=poppler_path,
                        )
                        text = clean_text(text)
                        ok = bool(text.strip())
                        warning = "PDF looked scanned; OCR applied automatically."

            elif ext in IMAGE_EXTS and enable_ocr:
                method = "ocr:image"
                text = _ocr_image_file(p, lang=ocr_lang)
                text = clean_text(text)
                ok = bool(text.strip())

            elif ext in CONVERTABLE_OFFICE_EXTS and enable_libreoffice and soffice:
                # Convert supported office formats via LibreOffice
                method = "convert:libreoffice->txt"
                with tempfile.TemporaryDirectory() as td:
                    converted = _convert_with_libreoffice(soffice, p, Path(td), target="txt:Text")
                    text = load_txt_like(converted)
                text = clean_text(text)
                ok = bool(text.strip())

            else:
                # Fallback: attempt conversion via LibreOffice for any other unknown extensions
                if enable_libreoffice and soffice:
                    try:
                        method = "convert:libreoffice->txt"
                        with tempfile.TemporaryDirectory() as td:
                            converted = _convert_with_libreoffice(soffice, p, Path(td), target="txt:Text")
                            text = load_txt_like(converted)
                        text = clean_text(text)
                        ok = bool(text.strip())
                        if not ok:
                            error = f"No text extracted for {ext}"
                    except Exception as ex:
                        ok = False
                        error = f"Unsupported or missing dependency for {ext}: {ex}"
                else:
                    ok = False
                    error = f"Unsupported or missing dependency for {ext}"

        except Exception as e:
            ok = False
            error = str(e)

        report.append(
            LoadReportEntry(
                source=str(p),
                ext=ext,
                ok=ok,
                method=method,
                chars=len(text),
                warning=warning,
                error=error,
            )
        )

        if ok:
            docs.append({"source": str(p), "text": text})

    return docs, report