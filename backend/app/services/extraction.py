import hashlib
import os
from typing import Tuple
import pdfplumber
from docx import Document as DocxDocument
import openpyxl


def sha256_file(path: str) -> str:
    hasher = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(8192), b""):
            hasher.update(chunk)
    return hasher.hexdigest()


def extract_text(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        with pdfplumber.open(path) as pdf:
            return "\n".join(page.extract_text() or "" for page in pdf.pages)
    if ext in {".docx"}:
        doc = DocxDocument(path)
        return "\n".join(paragraph.text for paragraph in doc.paragraphs)
    if ext in {".xlsx", ".xlsm"}:
        wb = openpyxl.load_workbook(path, data_only=True)
        parts = []
        for sheet in wb.worksheets:
            parts.append(f"Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = [str(cell) if cell is not None else "" for cell in row]
                parts.append("\t".join(values))
        return "\n".join(parts)
    with open(path, "r", encoding="utf-8", errors="ignore") as handle:
        return handle.read()


def save_extracted_text(root: str, doc_id: int, version: int, text: str) -> str:
    doc_dir = os.path.join(root, str(doc_id))
    os.makedirs(doc_dir, exist_ok=True)
    out_path = os.path.join(doc_dir, f"v{version}.txt")
    with open(out_path, "w", encoding="utf-8") as handle:
        handle.write(text)
    return out_path


def split_snippet(text: str, value: str, window: int = 120) -> Tuple[str, str]:
    idx = text.lower().find(value.lower()) if value else -1
    if idx == -1:
        snippet = text[: window * 2]
    else:
        start = max(idx - window, 0)
        end = min(idx + window, len(text))
        snippet = text[start:end]
    return snippet, "text"
